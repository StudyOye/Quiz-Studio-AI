import streamlit as st
import streamlit.components.v1 as components

# ------------------------------------------------------------------------------
# PWA ONE-CLICK INSTALL BUTTON COMPONENT
# ------------------------------------------------------------------------------
pwa_installer_code = """
<div id="pwa-install-wrapper" style="display:none; font-family: sans-serif; margin: 10px 0;">
    <button id="pwa-install-btn" style="
        width: 100%;
        background-color: #ff4b4b;
        color: #ffffff;
        border: none;
        padding: 12px 20px;
        font-size: 16px;
        font-weight: bold;
        border-radius: 8px;
        cursor: pointer;
        box-shadow: 0 4px 6px rgba(0,0,0,0.1);
        transition: background-color 0.2s ease;
    ">
        📲 Install App to Home Screen
    </button>
</div>

<script>
let deferredPrompt;

// 1. Capture the browser's install event
window.addEventListener('beforeinstallprompt', (e) => {
    // Prevent standard mini-infobar from showing
    e.preventDefault();
    deferredPrompt = e;
    
    // Unhide the custom Install button inside Streamlit
    document.getElementById('pwa-install-wrapper').style.display = 'block';
});

// 2. Trigger native prompt on click
document.getElementById('pwa-install-btn').addEventListener('click', async () => {
    if (deferredPrompt) {
        deferredPrompt.prompt();
        const { outcome } = await deferredPrompt.userChoice;
        if (outcome === 'accepted') {
            document.getElementById('pwa-install-wrapper').style.display = 'none';
        }
        deferredPrompt = null;
    }
});
</script>
"""

# Render the component in Streamlit layout
components.html(pwa_installer_code, height=70)

import os
import json
import io
import time
from typing import List, Dict, Any
from PIL import Image
import pypdf
import streamlit as st
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from docx import Document
from docx.shared import Pt as DocxPt, RGBColor as DocxRGBColor, Inches as DocxInches

# ------------------------------------------------------------------------------
# STREAMLIT PAGE CONFIGURATION
# ------------------------------------------------------------------------------
st.set_page_config(page_title="Quiz Studio AI", page_icon="🎯", layout="wide")


# ------------------------------------------------------------------------------
# GEMINI CLIENT INITIALIZATION (SECURE & CLOUD-READY)
# ------------------------------------------------------------------------------
def get_gemini_client():
    api_key = None
    # 1. Try fetching from Streamlit Secrets (for GitHub/Streamlit Cloud)
    try:
        if "GEMINI_API_KEY" in st.secrets:
            api_key = st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass

    # 2. Fall back to Environment Variables
    if not api_key:
        api_key = os.environ.get("GEMINI_API_KEY")

    if not api_key:
        st.error("🔑 Gemini API Key not found! Please add `GEMINI_API_KEY` to your Streamlit Secrets or Environment Variables.")
        st.stop()

    return genai.Client(api_key=api_key)


# ------------------------------------------------------------------------------
# HELPER FUNCTIONS & CONTENT EXTRACTION
# ------------------------------------------------------------------------------
def extract_text_from_pdfs(pdf_files) -> str:
    combined_text = ""
    for pdf_file in pdf_files:
        reader = pypdf.PdfReader(pdf_file)
        for page in reader.pages:
            combined_text += page.extract_text() or ""
    return combined_text


def generate_questions_with_ai(rounds_config: List[Dict], extracted_text: str, images: List[Any], target_language: str) -> Dict:
    client = get_gemini_client()

    prompt = f"""You are an expert quiz master building content for Quiz Studio AI.
Generate ALL output strictly in {target_language}.
Round Configuration: {json.dumps(rounds_config)}
Source Text: {extracted_text[:8000]}
Return STRICT JSON with keys: rounds -> round_name, questions -> id, question, options, answer, explanation."""

    contents = [prompt]
    for img in images:
        contents.append(Image.open(img))

    candidate_models = ["gemini-3.8-flash", "gemini-3.5-flash", "gemini-3.5-flash-lite"]
    last_err = None

    for model_name in candidate_models:
        for attempt in range(3):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents,
                    config=types.GenerateContentConfig(response_mime_type="application/json")
                )
                return json.loads(response.text)
            except Exception as e:
                last_err = e
                err_str = str(e).upper()
                if "429" in err_str or "RESOURCE_EXHAUSTED" in err_str:
                    time.sleep(6 * (attempt + 1))
                    continue
                else:
                    break
    raise last_err


# ------------------------------------------------------------------------------
# PRESENTATION & DOCUMENT BUILDERS
# ------------------------------------------------------------------------------
def create_pptx_deck(quiz_data: Dict, rounds_config: List[Dict], org_name: str, quiz_name: str, design_mode: str, text_size_mode: str) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    config_map = {r['name']: r for r in rounds_config}

    FONT_SCALES = {
        "small": {"title": 38, "round": 32, "question": 20, "option": 16, "answer": 32, "timer": 20, "hdr": 14},
        "medium": {"title": 48, "round": 40, "question": 24, "option": 18, "answer": 40, "timer": 22, "hdr": 16},
        "large": {"title": 56, "round": 46, "question": 28, "option": 22, "answer": 48, "timer": 26, "hdr": 18},
        "extra_large": {"title": 62, "round": 52, "question": 30, "option": 24, "answer": 54, "timer": 28, "hdr": 18}
    }
    fs = FONT_SCALES.get(text_size_mode, FONT_SCALES["medium"])

    BG_COLOR = RGBColor(255, 255, 255)
    TEXT_MAIN = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(51, 65, 85)
    ACCENT_PRIMARY = RGBColor(29, 78, 216)
    ANSWER_GREEN = RGBColor(21, 128, 61)
    TIMER_RED = RGBColor(225, 29, 72)
    KIDS_HEADER_BG, KIDS_ACCENT = RGBColor(254, 243, 199), RGBColor(217, 119, 6)

    def apply_design(slide):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BG_COLOR
        if design_mode == "minimalist":
            b1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
            b1.fill.solid(); b1.fill.fore_color.rgb = ACCENT_PRIMARY; b1.line.fill.background()
            b2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.38), Inches(13.333), Inches(0.12))
            b2.fill.solid(); b2.fill.fore_color.rgb = ACCENT_PRIMARY; b2.line.fill.background()
        elif design_mode == "heavy":
            hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(0.15), Inches(12.533), Inches(0.95))
            hdr.fill.solid(); hdr.fill.fore_color.rgb = KIDS_HEADER_BG
            hdr.line.color.rgb = KIDS_ACCENT; hdr.line.width = Pt(2)

    # Cover Slide
    slide = prs.slides.add_slide(blank_layout)
    apply_design(slide)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.0))
    tf = tb.text_frame; tf.word_wrap = True
    if org_name.strip():
        p_org = tf.paragraphs[0]; p_org.text = org_name.strip().upper()
        p_org.font.size = Pt(fs["option"]); p_org.font.bold = True; p_org.font.color.rgb = TEXT_MUTED; p_org.alignment = PP_ALIGN.CENTER
    p_title = tf.add_paragraph() if org_name.strip() else tf.paragraphs[0]
    p_title.text = quiz_name.strip().upper() if quiz_name.strip() else "QUIZ STUDIO AI"
    p_title.font.size = Pt(fs["title"]); p_title.font.bold = True; p_title.font.color.rgb = ACCENT_PRIMARY; p_title.alignment = PP_ALIGN.CENTER

    # Round & Question Slides
    for round_item in quiz_data.get("rounds", []):
        r_name = round_item["round_name"]
        r_config = config_map.get(r_name, {"timer": 30, "points": 10, "negative": 0})
        slide = prs.slides.add_slide(blank_layout)
        apply_design(slide)
        tb = slide.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(11.333), Inches(2.0))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]; p.text = f"ROUND: {r_name.upper()}"
        p.font.size = Pt(fs["round"]); p.font.bold = True; p.font.color.rgb = TEXT_MAIN; p.alignment = PP_ALIGN.CENTER

        for q_idx, q in enumerate(round_item["questions"]):
            slide_q = prs.slides.add_slide(blank_layout)
            apply_design(slide_q)
            
            tb_info = slide_q.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(7.2), Inches(0.8))
            tb_info.text_frame.word_wrap = True
            p_i = tb_info.text_frame.paragraphs[0]; p_i.text = f"{r_name} | Question {q_idx + 1}"
            p_i.font.size = Pt(fs["hdr"]); p_i.font.bold = True; p_i.font.color.rgb = ACCENT_PRIMARY if design_mode != "heavy" else KIDS_ACCENT
            
            tb_timer_lbl = slide_q.shapes.add_textbox(Inches(8.0), Inches(0.25), Inches(2.8), Inches(0.8))
            p_tl = tb_timer_lbl.text_frame.paragraphs[0]; p_tl.text = "⏱️ Time Limit (s):"
            p_tl.font.size = Pt(fs["hdr"] - 2); p_tl.font.bold = True; p_tl.font.color.rgb = TEXT_MUTED; p_tl.alignment = PP_ALIGN.RIGHT
            
            tb_timer_num = slide_q.shapes.add_textbox(Inches(11.0), Inches(0.22), Inches(1.5), Inches(0.8))
            p_tn = tb_timer_num.text_frame.paragraphs[0]; p_tn.text = str(r_config['timer'])
            p_tn.font.size = Pt(fs["timer"] + 2); p_tn.font.bold = True; p_tn.font.color.rgb = TIMER_RED; p_tn.alignment = PP_ALIGN.CENTER
            
            tb_q = slide_q.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(2.2))
            tb_q.text_frame.word_wrap = True
            p_q = tb_q.text_frame.paragraphs[0]; p_q.text = q["question"]
            p_q.font.size = Pt(fs["question"]); p_q.font.bold = True; p_q.font.color.rgb = TEXT_MAIN
            
            options_list = q.get("options", [])
            if options_list:
                top_pos = 3.7
                opt_gap = 0.8 if text_size_mode != "extra_large" else 0.7
                for opt_idx, opt in enumerate(options_list):
                    tb_opt = slide_q.shapes.add_textbox(Inches(1.0), Inches(top_pos), Inches(11.3), Inches(0.65))
                    tb_opt.text_frame.word_wrap = True
                    p_opt = tb_opt.text_frame.paragraphs[0]
                    p_opt.text = f"{chr(65 + opt_idx)}. {opt}"
                    p_opt.font.size = Pt(fs["option"]); p_opt.font.color.rgb = TEXT_MAIN
                    top_pos += opt_gap

            # Answer Slide
            slide_a = prs.slides.add_slide(blank_layout)
            apply_design(slide_a)
            tb_a_info = slide_a.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(11.7), Inches(0.8))
            tb_a_info.text_frame.word_wrap = True
            p_ai = tb_a_info.text_frame.paragraphs[0]; p_ai.text = f"{r_name} | Question {q_idx + 1} - CORRECT ANSWER"
            p_ai.font.size = Pt(fs["hdr"]); p_ai.font.bold = True; p_ai.font.color.rgb = TEXT_MUTED
            
            tb_ans = slide_a.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.333), Inches(3.5))
            tb_ans.text_frame.word_wrap = True
            p_ans_lbl = tb_ans.text_frame.paragraphs[0]; p_ans_lbl.text = "CORRECT ANSWER:"
            p_ans_lbl.font.size = Pt(fs["option"]); p_ans_lbl.font.bold = True; p_ans_lbl.font.color.rgb = TEXT_MUTED; p_ans_lbl.alignment = PP_ALIGN.CENTER
            
            p_ans_val = tb_ans.text_frame.add_paragraph()
            p_ans_val.text = q["answer"]; p_ans_val.font.size = Pt(fs["answer"])
            p_ans_val.font.bold = True; p_ans_val.font.color.rgb = ANSWER_GREEN; p_ans_val.alignment = PP_ALIGN.CENTER

    # Thank You Slide
    slide_ty = prs.slides.add_slide(blank_layout)
    apply_design(slide_ty)
    tb_ty = slide_ty.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.333), Inches(2.5))
    tb_ty.text_frame.word_wrap = True
    p_ty = tb_ty.text_frame.paragraphs[0]; p_ty.text = "THANK YOU!"
    p_ty.font.size = Pt(fs["title"]); p_ty.font.bold = True; p_ty.font.color.rgb = ACCENT_PRIMARY; p_ty.alignment = PP_ALIGN.CENTER
    
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def create_rulebook_pdf(rounds_config: List[Dict]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet(); story = []
    
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=22, textColor=colors.HexColor("#1D4ED8"))
    story.append(Paragraph("Quiz Studio AI - Official Rulebook", title_style)); story.append(Spacer(1, 10))
    table_data = [["Round Name", "Options", "Timer", "Points (+)", "Negative (-)", "Buzzer"]]
    for r in rounds_config:
        table_data.append([
            r['name'],
            f"{r['num_options']} Choices" if r['num_options'] > 0 else 'Direct Q&A',
            f"{r['timer']}s",
            f"+{r['points']}",
            f"-{r['negative']}",
            'Enabled' if r['buzzer'] else 'Disabled'
        ])
    table = Table(table_data, colWidths=[130, 90, 60, 70, 70, 70])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1D4ED8")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('GRID', (0,0), (-1,-1), 1, colors.HexColor("#CBD5E1")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE')
    ]))
    story.append(table); doc.build(story); buffer.seek(0)
    return buffer.getvalue()


def create_qa_master_pdf(quiz_data: Dict) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet(); story = []
    
    title_style = ParagraphStyle('QATitle', parent=styles['Heading1'], fontSize=20, leading=24, textColor=colors.HexColor("#1D4ED8"))
    round_style = ParagraphStyle('QARound', parent=styles['Heading2'], fontSize=15, leading=19, textColor=colors.HexColor("#0F172A"), spaceBefore=10)
    q_style = ParagraphStyle('QAQuestion', parent=styles['Normal'], fontSize=11, leading=15, textColor=colors.HexColor("#1E293B"))
    ans_style = ParagraphStyle('QAAnswer', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor("#166534"))
    exp_style = ParagraphStyle('QAExplanation', parent=styles['Normal'], fontSize=9, leading=13, textColor=colors.HexColor("#475569"))

    story.append(Paragraph("Quiz Studio AI - Master Answer Key", title_style)); story.append(Spacer(1, 10))
    for r in quiz_data.get('rounds', []):
        story.append(Paragraph(f"📖 {r['round_name']}", round_style))
        story.append(Spacer(1, 4))
        for idx, q in enumerate(r['questions']):
            opts = f"<br/><i>Options:</i> {', '.join(q['options'])}" if q.get('options') else ''
            story.append(Paragraph(f"Q{idx+1}: {q['question']}{opts}", q_style))
            story.append(Paragraph(f"<b>Correct Answer:</b> {q['answer']}", ans_style))
            story.append(Paragraph(f"<i>Explanation:</i> {q['explanation']}", exp_style))
            story.append(Spacer(1, 6))
    doc.build(story); buffer.seek(0)
    return buffer.getvalue()


def create_qa_master_docx(quiz_data: Dict) -> bytes:
    doc = Document()
    
    # Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Quiz Studio AI - Master Answer Key")
    r_title.font.size = DocxPt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = DocxRGBColor(29, 78, 216)
    
    for r in quiz_data.get('rounds', []):
        p_rnd = doc.add_paragraph()
        r_rnd = p_rnd.add_run(f"📖 {r['round_name']}")
        r_rnd.font.size = DocxPt(14)
        r_rnd.font.bold = True
        r_rnd.font.color.rgb = DocxRGBColor(15, 23, 42)
        
        for idx, q in enumerate(r['questions']):
            p_q = doc.add_paragraph()
            r_q = p_q.add_run(f"Q{idx+1}: {q['question']}")
            r_q.font.bold = True
            r_q.font.size = DocxPt(11)
            
            if q.get('options'):
                p_opts = doc.add_paragraph()
                r_opts = p_opts.add_run(f"Options: {', '.join(q['options'])}")
                r_opts.font.italic = True
                r_opts.font.size = DocxPt(10)
                
            p_ans = doc.add_paragraph()
            p_ans.add_run("Correct Answer: ").font.bold = True
            r_ans_val = p_ans.add_run(q['answer'])
            r_ans_val.font.color.rgb = DocxRGBColor(22, 101, 52)
            
            p_exp = doc.add_paragraph()
            r_exp = p_exp.add_run(f"Explanation: {q['explanation']}")
            r_exp.font.italic = True
            r_exp.font.size = DocxPt(9.5)
            r_exp.font.color.rgb = DocxRGBColor(71, 85, 105)
            
            doc.add_paragraph().paragraph_format.space_after = DocxPt(4)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------------------
# STREAMLIT UI LAYOUT & CONTROLS
# ------------------------------------------------------------------------------
st.title("🎯 Quiz Studio AI")

target_language = st.sidebar.selectbox("Select Output PPT Language:", options=["English", "Punjabi", "Hindi"], index=0)
org_name = st.sidebar.text_input("Organisation Name", value="GMSSS Dhrangwala")
quiz_name = st.sidebar.text_input("Quiz Name / Title", value="Shiksha-Hack 2026 IT Quiz")
uploaded_pdfs = st.sidebar.file_uploader("Upload Source PDFs", type=["pdf"], accept_multiple_files=True)
uploaded_imgs = st.sidebar.file_uploader("Upload Source JPEGs", type=["jpeg", "jpg"], accept_multiple_files=True)
design_mode = st.sidebar.selectbox("PPT Design Mode:", options=["no_design", "minimalist", "heavy"], index=0)
text_size_mode = st.sidebar.selectbox("Text Size Option:", options=["small", "medium", "large", "extra_large"], index=1)

num_rounds = st.number_input("Number of Rounds", min_value=1, max_value=6, value=2)
rounds_config = []
cols = st.columns(num_rounds)
for i in range(num_rounds):
    with cols[i]:
        st.subheader(f"Round {i+1}")
        r_name = st.text_input("Name", value=f"Round {i+1}", key=f"rname_{i}")
        q_count = st.number_input("Questions", min_value=1, max_value=10, value=3, key=f"qcnt_{i}")
        num_options = st.selectbox("Options per Question:", options=[0, 2, 3, 4], index=3, key=f"opts_{i}")
        timer = st.number_input("Timer (sec)", min_value=5, max_value=120, value=30, key=f"timer_{i}")
        points = st.number_input("Points (+)", min_value=1, max_value=50, value=10, key=f"pts_{i}")
        negative = st.number_input("Negative (-)", min_value=0, max_value=50, value=2, key=f"neg_{i}")
        buzzer = st.checkbox("Buzzer Mode", value=False, key=f"buzz_{i}")
        
        rounds_config.append({
            'name': r_name,
            'count': q_count,
            'num_options': num_options,
            'timer': timer,
            'points': points,
            'negative': negative,
            'buzzer': buzzer
        })

if st.button("🚀 Generate Presentation & Printable Documents", type="primary"):
    if not uploaded_pdfs and not uploaded_imgs:
        st.warning("Please upload at least one PDF or JPEG file.")
    else:
        with st.spinner(f"Extracting content and querying Gemini AI in {target_language}..."):
            extracted_text = extract_text_from_pdfs(uploaded_pdfs) if uploaded_pdfs else ""
            quiz_data = generate_questions_with_ai(rounds_config, extracted_text, uploaded_imgs or [], target_language)
            st.session_state['quiz_data'] = quiz_data
            st.success("Quiz bundle generated successfully!")

if 'quiz_data' in st.session_state:
    quiz_data = st.session_state['quiz_data']
    col1, col2, col3 = st.columns(3)
    
    with col1:
        st.download_button(
            "📊 Slide Deck (.pptx)",
            create_pptx_deck(quiz_data, rounds_config, org_name, quiz_name, design_mode, text_size_mode),
            "Quiz_Deck.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            use_container_width=True
        )
    with col2:
        st.download_button(
            "📑 Rulebook PDF (.pdf)",
            create_rulebook_pdf(rounds_config),
            "Rulebook.pdf",
            mime="application/pdf",
            use_container_width=True
        )
    with col3:
        # Dynamic QA Key format based on selected language
        if target_language == "English":
            st.download_button(
                "🖨️ Master QA Key (.pdf)",
                create_qa_master_pdf(quiz_data),
                "Master_QA_Key.pdf",
                mime="application/pdf",
                use_container_width=True
            )
        else:
            st.download_button(
                "🖨️ Master QA Key (.docx)",
                create_qa_master_docx(quiz_data),
                "Master_QA_Key.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

    st.markdown("---")
    st.subheader("💡 Microsoft PowerPoint VBA Countdown Script")
    st.caption("Open MS PowerPoint Desktop -> Press Alt + F11 -> Insert -> Module -> Paste Code:")
    st.code("""
#If VBA7 Then
    Private Declare PtrSafe Sub Sleep Lib "kernel32" (ByVal dwMilliseconds As Long)
#Else
    Private Declare Sub Sleep Lib "kernel32" (ByVal dwMilliseconds As Long)
#End If

Sub StartQuestionTimer()
    Dim shp As Shape
    Dim slideObj As Slide
    Dim sec As Integer
    Dim startVal As Integer
    
    On Error Resume Next
    Set slideObj = Application.ActivePresentation.SlideShowWindow.View.Slide
    
    For Each shp In slideObj.Shapes
        If shp.HasTextFrame Then
            If IsNumeric(Trim(shp.TextFrame.TextRange.Text)) Then
                startVal = CInt(Trim(shp.TextFrame.TextRange.Text))
                If startVal > 0 And startVal <= 300 Then
                    For sec = startVal To 0 Step -1
                        shp.TextFrame.TextRange.Text = CStr(sec)
                        DoEvents
                        Sleep 1000
                    Next sec
                    Exit Sub
                End If
            End If
        End If
    Next shp
End Sub
""", language="vb")
